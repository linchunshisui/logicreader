
from docx import Document
from docx.shared import Pt
doc = Document()
doc.add_heading('LogicReader DOCX Fixture', level=1)
doc.add_paragraph('This document exercises the DOCX pipeline: headings, paragraphs, a table and an image.')
doc.add_heading('1 Claims and evidence', level=2)
doc.add_paragraph('Attention replaces recurrence with direct pairwise interaction, which improves both accuracy and training throughput.')
doc.add_paragraph('The proposed method reaches 28.4 BLEU on the English to German task.')
doc.add_heading('1.1 Evidence table', level=3)
table = doc.add_table(rows=3, cols=3)
data = [['Model', 'BLEU', 'Parallel'], ['RNN baseline', '26.3', 'no'], ['Attention', '28.4', 'yes']]
for r, row in enumerate(data):
    for c, value in enumerate(row):
        table.cell(r, c).text = value
doc.add_heading('2 Figure', level=2)
doc.add_picture(r'D:\逻辑阅读器\tests\fixtures\chart.png')
doc.add_paragraph('Figure 1: quarterly revenue by product line.')
doc.add_heading('3 Conclusion', level=2)
doc.add_paragraph('Attention mechanisms keep the model parallel while improving quality.')
doc.save(r'D:\逻辑阅读器\tests\fixtures\report.docx')

from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = 'Revenue'
ws.append(['Quarter', 'Product A', 'Product B', 'Total'])
for i, q in enumerate(['Q1', 'Q2', 'Q3', 'Q4']):
    a = 120 + i * 15
    b = 90 + i * 22
    ws.append([q, a, b, None])
    ws.cell(row=i + 2, column=4).value = f'=SUM(B{i+2}:C{i+2})'
ws2 = wb.create_sheet('Notes')
ws2.append(['note', 'value'])
ws2.append(['reviewed', 'yes'])
wb.save(r'D:\逻辑阅读器\tests\fixtures\finance.xlsx')
print('docx + xlsx ok')
